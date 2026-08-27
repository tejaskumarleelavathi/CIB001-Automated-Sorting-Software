"""
cib/deployment/generate_plate_config.py — build/update one station's sheet
in a multi-station plate config workbook, from 3 measured reference wells.

Why this exists (2026-08-20, see DEV-17 in prd_updated.md)
------------------------------------------------------------
There are 4 physical plate stations (nests). Per the PRD's existing
"station map" design (FR-SETUP-01), a station is either populated with a
plate or not — but SOMETHING still has to supply the well geometry (every
well's X/Y in robot coordinates) for whichever plate is seated in each
station. This script is that something: same 3-point grid math as the
original single-station hardware_tests/plate_from_coords.py, but reusable
per-station and CLI-driven so it isn't a one-off script tied to specific
hardcoded coordinates.

One file, one sheet per station ("Station 1".."Station N") — same
well/x/y/... schema every single-station plate config already used, just
one workbook holding all of them (see cib/config/loader.py's
load_all_stations()). Run this once per station as you get that station's
measured A1/A12/H1 corners; re-running for a station you've already done
replaces just that station's sheet, leaving the others untouched.

U (wrist angle) is DELIBERATELY NOT part of this file or its output — it
lives in config/machine_config.yaml's plate_stations: section instead. A
station's U is a fixed property of where that nest physically is, not
something that should ever be edited by whoever is filling in well
coordinates. This script will remind you to check/update it, not write it.

USAGE
  python -m cib.deployment.generate_plate_config \
      --station 1 \
      --a1  189.955,287.602 \
      --a12 189.955,188.602 \
      --h1  252.955,287.602 \
      --out config/plate_configs/plate_stations.xlsx

  Optional: --rows 8 --cols 12 (default: standard 96-well, 8 rows A-H x 12 cols)

  NOTE — left-side stations have negative X coordinates. argparse reads a
  bare "--a1 -229.636,269.400" as an unknown flag, not a value. Use the
  --flag=value form instead: --a1=-229.636,269.400
"""

from __future__ import annotations

import argparse
import os

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["well", "x", "y", "aspirate_volume_ul", "aspirate_speed_ul_per_s",
           "dispense_volume_ul", "dispense_speed_ul_per_s"]


def _parse_xy(s: str) -> tuple[float, float]:
    x_str, y_str = s.split(",")
    return float(x_str), float(y_str)


def compute_grid(a1: tuple, a12: tuple, h1: tuple, rows: int, cols: int) -> dict[str, tuple[float, float]]:
    """
    Same math as hardware_tests/plate_from_coords.py / teach_plate_wells.py:
      col_vector = (A{cols} - A1) / (cols - 1)   -- per-column step (A1 -> A{cols} direction)
      row_vector = ({row H} 1  - A1) / (rows - 1) -- per-row step    (A1 -> last-row-1 direction)
      well(r, c) = A1 + c*col_vector + r*row_vector
    """
    a1_v = np.array(a1, dtype=float)
    a12_v = np.array(a12, dtype=float)
    h1_v = np.array(h1, dtype=float)

    col_vec = (a12_v - a1_v) / (cols - 1)
    row_vec = (h1_v - a1_v) / (rows - 1)

    print(f"  Column vector (per well, A1->A{cols} direction): {col_vec} mm  "
          f"|v|={np.linalg.norm(col_vec):.3f}mm")
    print(f"  Row vector    (per well, row A->{chr(ord('A') + rows - 1)} direction): {row_vec} mm  "
          f"|v|={np.linalg.norm(row_vec):.3f}mm")
    print("  (nominal 96-well pitch is 9.00mm — compare against the above)")

    wells: dict[str, tuple[float, float]] = {}
    for r in range(rows):
        for c in range(cols):
            pos = a1_v + c * col_vec + r * row_vec
            label = f"{chr(ord('A') + r)}{c + 1}"
            wells[label] = (float(pos[0]), float(pos[1]))
    return wells


def write_station_sheet(out_path: str, station: int, wells: dict[str, tuple[float, float]],
                         rows: int, cols: int) -> None:
    """
    Write/replace the "Station <N>" sheet in out_path, creating the workbook
    if it doesn't exist yet. Other stations' sheets are left untouched.
    """
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    sheet_name = f"Station {station}"

    if os.path.exists(out_path):
        wb = openpyxl.load_workbook(out_path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        # A brand-new workbook always starts with one default empty sheet
        # ("Sheet") — drop it once we're about to add real content, but
        # only if it's still empty and unused.
        if "Sheet" in wb.sheetnames and wb["Sheet"].max_row <= 1 and wb["Sheet"].max_column <= 1:
            del wb["Sheet"]
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # drop the default sheet, we name our own below

    ws = wb.create_sheet(sheet_name)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    row_idx = 2
    for r in range(rows):
        for c in range(cols):
            label = f"{chr(ord('A') + r)}{c + 1}"
            x, y = wells[label]
            ws.cell(row=row_idx, column=1, value=label).border = border
            ws.cell(row=row_idx, column=2, value=round(x, 3)).border = border
            ws.cell(row=row_idx, column=3, value=round(y, 3)).border = border
            for col in (4, 5, 6, 7):
                ws.cell(row=row_idx, column=col).border = border
            row_idx += 1

    col_widths = [10, 12, 12, 20, 22, 20, 22]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # Keep station sheets in numeric order regardless of what order they
    # were generated/updated in.
    order = sorted(wb.sheetnames, key=lambda n: (0, int(n.split(" ")[1])) if n.startswith("Station ") else (1, n))
    wb._sheets = [wb[n] for n in order]

    wb.save(out_path)
    print(f"\n  Saved: {out_path}  (sheet '{sheet_name}', {len(wells)} wells)")
    print(f"  Workbook now has sheets: {wb.sheetnames}")


def main() -> None:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--station", type=int, required=True, help="Station number (1-4)")
    ap.add_argument("--a1", type=str, required=True, help="A1 coordinate as 'x,y' mm")
    ap.add_argument("--a12", type=str, required=True,
                     help="Last well of row A (e.g. A12 for a 96-well plate) as 'x,y' mm")
    ap.add_argument("--h1", type=str, required=True,
                     help="First well of the last row (e.g. H1 for a 96-well plate) as 'x,y' mm")
    ap.add_argument("--rows", type=int, default=8, help="Number of rows (default 8, A-H)")
    ap.add_argument("--cols", type=int, default=12, help="Number of columns (default 12)")
    ap.add_argument("--out", type=str, required=True,
                     help="Path to the multi-station workbook (created if it doesn't exist)")
    args = ap.parse_args()

    a1 = _parse_xy(args.a1)
    a12 = _parse_xy(args.a12)
    h1 = _parse_xy(args.h1)

    print("=" * 70)
    print(f"  Generating Station {args.station} ({args.rows} rows x {args.cols} cols)")
    print("=" * 70)

    wells = compute_grid(a1, a12, h1, args.rows, args.cols)

    print("\n  Computed wells (first row):")
    for c in range(args.cols):
        label = f"A{c + 1}"
        wx, wy = wells[label]
        print(f"    {label:<4} ({wx:.2f}, {wy:.2f}) mm")

    write_station_sheet(args.out, args.station, wells, args.rows, args.cols)

    print(f"\n  NOTE: U (wrist angle) is NOT stored in this file — it lives in")
    print(f"  config/machine_config.yaml's plate_stations: section. Check/update")
    print(f"  the u_deg for station {args.station} there if it isn't set yet.")


if __name__ == "__main__":
    main()
