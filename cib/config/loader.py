"""
cib/config/loader.py — Load and parse machine config (YAML) and plate config (Excel).

Usage:
    from cib.config.loader import load_machine_config, load_plate_config, load_all_stations

    machine = load_machine_config()                      # reads default path
    machine = load_machine_config(Path("custom.yaml"))  # or a custom path

    plate = load_plate_config(Path(r"C:\LabAutomation\config\plate_configs\96well.xlsx"))

    # Multi-station workbook (one "Station N" sheet per physical nest):
    stations = load_all_stations(Path("config/plate_configs/plate_stations.xlsx"))
    # -> {1: PlateConfig(...), 2: PlateConfig(...), ...}

Both functions return typed dataclass objects. Validators run automatically —
ConfigValidationError is raised if anything is wrong.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

import yaml
import openpyxl

from cib.config.validator import ConfigValidationError, validate_machine_config, validate_plate_rows


DEFAULT_MACHINE_CONFIG = Path(r"C:\LabAutomation\config\machine_config.yaml")


# ── Plate config dataclasses ──────────────────────────────────────────────────

@dataclass
class WellConfig:
    """One row from the plate config Excel — one well's definition."""
    well:                   str
    x:                      float
    y:                      float
    aspirate_volume_ul:     Optional[float] = None
    aspirate_speed_ul_per_s: Optional[float] = None
    dispense_volume_ul:     Optional[float] = None
    dispense_speed_ul_per_s: Optional[float] = None


@dataclass
class PlateConfig:
    """
    Loaded plate config. wells is an ordered list — order defines fill sequence.
    First well in the list = first well the robot fills.
    """
    source_file: str
    wells:       List[WellConfig]

    def well_names(self) -> List[str]:
        return [w.well for w in self.wells]

    def well_by_name(self, name: str) -> Optional[WellConfig]:
        name_upper = name.upper()
        for w in self.wells:
            if w.well.upper() == name_upper:
                return w
        return None

    def index_of(self, well_name: str) -> Optional[int]:
        """Return the 0-based index of a well in the fill order, or None if not found."""
        name_upper = well_name.upper()
        for i, w in enumerate(self.wells):
            if w.well.upper() == name_upper:
                return i
        return None

    def wells_from(self, start_well: str) -> List[WellConfig]:
        """
        Return all wells starting from start_well (inclusive) in fill order.
        Used by orchestrator when operator specifies a starting well.
        Raises ConfigValidationError if start_well is not in this plate config.
        """
        idx = self.index_of(start_well)
        if idx is None:
            raise ConfigValidationError(
                f"Starting well '{start_well}' is not in plate config '{self.source_file}'. "
                f"Available wells: {self.well_names()}"
            )
        return self.wells[idx:]


# ── Machine config dataclasses ────────────────────────────────────────────────

@dataclass
class RobotConfig:
    ip:                 str
    port:               int
    connect_timeout_s:  float
    command_timeout_s:  float
    # Quadrant-based U (wrist) angle selection — all three optional. If any is
    # None, quadrant U-selection is disabled (MOVE_IMAGE keeps holding whatever
    # U the robot is currently at — pre-2026-08-05 behavior).
    quadrant_split_y_mm: Optional[float] = None
    u_top_deg:           Optional[float] = None
    u_bottom_deg:        Optional[float] = None


@dataclass
class RobotRemoteConfig:
    """
    Epson RC+ Remote Ethernet — controller-level lifecycle control (Login,
    GetStatus, Reset, SetMotorsOn/Off, Start, Stop), separate from the
    robot's own PICK/DROP/MOVE_IMAGE protocol (see RobotConfig above).
    Requires a one-time setup in RC+: Control Device = Remote Ethernet
    (System Configuration -> Controller -> Configuration), plus a port and
    password on the System Configuration -> Remote Control -> Ethernet page.
    enabled=False disables the whole feature (no-op) so a controller without
    that setup can still run normally.
    """
    ip:                str
    port:              int
    password:          str
    connect_timeout_s: float
    command_timeout_s: float
    enabled:           bool = False


@dataclass
class PLCConfig:
    ip:              str
    port:            int
    poll_interval_ms: int
    unit_id:         int
    registers:       dict
    coils:           dict


@dataclass
class FlexiConfig:
    ip:              str
    port:            int
    poll_interval_ms: int
    unit_id:         int
    registers:       dict


@dataclass
class CameraConfig:
    ip:                      str
    serial:                  str
    exposure_us:             float
    gain_db:                 float
    pixel_format:            str = "Mono8"
    ip_original:             str = ""      # for EoA camera force_ip
    backlight_on_exposure_us: Optional[float] = None  # if set, use this shorter exposure for the backlight-ON frame


@dataclass
class BacklightConfig:
    ip:                str
    port:              int
    current_ma:        int
    connect_timeout_s: float


@dataclass
class RingLightConfig:
    robot_output_pin: int


@dataclass
class PipetteConfig:
    interface:              str
    ip:                     str
    port:                   int
    aspirate_volume_ul:     float
    aspirate_speed_ul_per_s: float
    dispense_volume_ul:     float
    dispense_speed_ul_per_s: float


@dataclass
class PathsConfig:
    db:              str
    plate_configs:   str
    runs_output:     str
    logs:            str
    overhead_images: str = "output/overhead_images"
    # Raw (pre-vision, pre-annotation) overhead captures -- separate from
    # overhead_images above, which only ever holds the annotated copy.
    # 2026-09-08, operator request.
    unprocessed_overhead_images: str = "output/unprocessed_overhead"


@dataclass
class VisionConfig:
    detection_min_area_px:  int
    detection_max_area_px:  int
    backlight_lead_time_ms: int
    overhead_save_images:   bool  = True
    overhead_save_format:   str   = "png"
    eoa_save_format:        str   = "tiff"
    centroid_snap_mm:       float = 0.8


@dataclass
class LoggingConfig:
    level:          str
    retention_days: int


@dataclass
class PlateStationsConfig:
    """
    Per-side home/quadrant U angles plus the station->side mapping, for the
    multi-station Lefty/Righty sorting cycle. Optional -- absent from
    machine_config.yaml entirely means the single-plate, Righty-only cycle
    (MachineConfig.plate_stations is None).
    """
    right_home_u:    float
    left_home_u:     float
    righty_top_u:    float
    righty_bottom_u: float
    lefty_top_u:     float
    lefty_bottom_u:  float
    stations:        Dict[int, str]   # station number -> "right" | "left"

    def side_of(self, station: int) -> str:
        side = self.stations.get(station)
        if side is None:
            raise ConfigValidationError(
                f"Station {station} has no 'side' configured in "
                f"machine_config.yaml's plate_stations.stations section."
            )
        return side


@dataclass
class MachineConfig:
    robot:           RobotConfig
    robot_remote:    RobotRemoteConfig
    plc:             PLCConfig
    flexi:           FlexiConfig
    overhead_camera: CameraConfig
    eoa_camera:      CameraConfig
    backlight:       BacklightConfig
    ring_light:      RingLightConfig
    pipette:         PipetteConfig
    paths:           PathsConfig
    vision:          VisionConfig
    logging:         LoggingConfig
    plate_stations:  Optional[PlateStationsConfig] = None


# ── Loaders ───────────────────────────────────────────────────────────────────

def load_machine_config(path: Path = DEFAULT_MACHINE_CONFIG) -> MachineConfig:
    """
    Load and validate machine_config.yaml.
    Raises FileNotFoundError if the file doesn't exist.
    Raises ConfigValidationError if any value is invalid.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"Machine config not found: {path}\n"
            f"Copy config/machine_config.yaml from the repo to {path}"
        )

    with open(path, "r", encoding="utf-8") as f:
        raw = yaml.safe_load(f)

    if not isinstance(raw, dict):
        raise ConfigValidationError(
            f"machine_config.yaml appears to be empty or malformed: {path}"
        )

    validate_machine_config(raw)
    return _parse_machine_config(raw)


def load_plate_config(path: Path) -> PlateConfig:
    """
    Load and validate a plate config Excel file (.xlsx).
    Row order in the file defines the well fill order.
    Raises FileNotFoundError if the file doesn't exist.
    Raises ConfigValidationError if any value is invalid.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Plate config not found: {path}")

    rows = _read_excel(path)
    validate_plate_rows(rows, source=path.name)
    wells = [_parse_well_row(r) for r in rows]
    return PlateConfig(source_file=path.name, wells=wells)


def load_all_stations(path: Path) -> Dict[int, PlateConfig]:
    """
    Load a multi-sheet plate-stations workbook (one "Station N" sheet per
    physical nest, e.g. config/plate_configs/plate_stations.xlsx) into
    {station_number: PlateConfig}. Same well/x/y/... column schema as
    load_plate_config's single-sheet files -- one sheet per station instead
    of one file per plate. Sheets not named "Station <int>" are ignored.
    Raises ConfigValidationError if no station sheets with data are found,
    or if any sheet's rows fail the normal plate-row validation.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Plate stations config not found: {path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    stations: Dict[int, PlateConfig] = {}
    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("Station "):
            continue
        try:
            station_num = int(sheet_name.split(" ")[1])
        except (IndexError, ValueError):
            continue
        rows = _rows_from_sheet(wb[sheet_name])
        if not rows:
            continue
        source = f"{path.name}:{sheet_name}"
        validate_plate_rows(rows, source=source)
        wells = [_parse_well_row(r) for r in rows]
        stations[station_num] = PlateConfig(source_file=source, wells=wells)
    wb.close()

    if not stations:
        raise ConfigValidationError(
            f"{path.name}: no 'Station N' sheets found with data."
        )
    return stations


def list_plate_configs(config_dir: Path) -> List[Path]:
    """
    Return all .xlsx files in the plate config directory, sorted by name.
    Used by the UI to populate the config selection dropdown.
    """
    config_dir = Path(config_dir)
    if not config_dir.exists():
        return []
    return sorted(config_dir.glob("*.xlsx"))


# ── Internal parsers ──────────────────────────────────────────────────────────

def _read_excel(path: Path) -> list[dict]:
    """
    Read an Excel file's active sheet and return rows as a list of dicts
    with normalised keys. First row is treated as the header.
    """
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    result = _rows_from_sheet(wb.active)
    wb.close()
    return result


def _rows_from_sheet(ws) -> list[dict]:
    """
    Shared by _read_excel (single-plate files) and load_all_stations
    (multi-sheet plate-stations workbook). First row is treated as the
    header, normalised to lowercase stripped strings. Blank rows skipped.
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    result = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        result.append({headers[i]: row[i] for i in range(len(headers))})

    return result


def _parse_well_row(row: dict) -> WellConfig:
    def _opt_float(key: str) -> Optional[float]:
        val = row.get(key)
        if val is None or str(val).strip() == "":
            return None
        return float(val)

    return WellConfig(
        well=str(row["well"]).strip().upper(),
        x=float(row["x"]),
        y=float(row["y"]),
        aspirate_volume_ul=_opt_float("aspirate_volume_ul"),
        aspirate_speed_ul_per_s=_opt_float("aspirate_speed_ul_per_s"),
        dispense_volume_ul=_opt_float("dispense_volume_ul"),
        dispense_speed_ul_per_s=_opt_float("dispense_speed_ul_per_s"),
    )


def _parse_machine_config(raw: dict) -> MachineConfig:
    r = raw["robot"]
    # Optional section — Remote Ethernet requires a one-time RC+ setup that
    # doesn't exist on every controller. .get() with defaults so older
    # machine_config.yaml files without this section still load.
    rr = raw.get("robot_remote", {})
    p = raw["plc"]
    fl = raw["flexi"]
    oc = raw["overhead_camera"]
    ec = raw["eoa_camera"]
    bl = raw["backlight"]
    rl = raw.get("ring_light", {})
    pp = raw["pipette"]
    pa = raw["paths"]
    vi = raw["vision"]
    lo = raw["logging"]
    ps = raw.get("plate_stations")

    return MachineConfig(
        robot=RobotConfig(
            ip=r["ip"],
            port=int(r["port"]),
            connect_timeout_s=float(r["connect_timeout_s"]),
            command_timeout_s=float(r["command_timeout_s"]),
            quadrant_split_y_mm=(
                float(r["quadrant_split_y_mm"]) if r.get("quadrant_split_y_mm") is not None else None
            ),
            u_top_deg=(
                float(r["u_top_deg"]) if r.get("u_top_deg") is not None else None
            ),
            u_bottom_deg=(
                float(r["u_bottom_deg"]) if r.get("u_bottom_deg") is not None else None
            ),
        ),
        robot_remote=RobotRemoteConfig(
            ip=rr.get("ip", r["ip"]),   # same controller as robot: by default
            port=int(rr.get("port", 5000)),
            password=str(rr.get("password", "")),
            connect_timeout_s=float(rr.get("connect_timeout_s", 5.0)),
            command_timeout_s=float(rr.get("command_timeout_s", 10.0)),
            enabled=bool(rr.get("enabled", False)),
        ),
        plc=PLCConfig(
            ip=p["ip"],
            port=int(p["port"]),
            poll_interval_ms=int(p["poll_interval_ms"]),
            unit_id=int(p.get("unit_id", 1)),
            registers=p.get("registers", {}),
            coils=p.get("coils", {}),
        ),
        flexi=FlexiConfig(
            ip=fl["ip"],
            port=int(fl["port"]),
            poll_interval_ms=int(fl["poll_interval_ms"]),
            unit_id=int(fl.get("unit_id", 1)),
            registers=fl.get("registers", {}),
        ),
        overhead_camera=CameraConfig(
            ip=oc["ip"],
            serial=oc.get("serial", ""),
            exposure_us=float(oc.get("exposure_us", 2000)),
            gain_db=float(oc.get("gain_db", 0.0)),
            pixel_format=oc.get("pixel_format", "Mono8"),
        ),
        eoa_camera=CameraConfig(
            ip=ec["ip"],
            serial=ec.get("serial", ""),
            exposure_us=float(ec.get("exposure_us", 120000)),
            gain_db=float(ec.get("gain_db", 0.0)),
            pixel_format=ec.get("pixel_format", "BGR8"),
            ip_original=ec.get("ip_original", ""),
            backlight_on_exposure_us=(
                float(ec["backlight_on_exposure_us"])
                if ec.get("backlight_on_exposure_us") is not None else None
            ),
        ),
        backlight=BacklightConfig(
            ip=bl["ip"],
            port=int(bl["port"]),
            current_ma=int(bl["current_ma"]),
            connect_timeout_s=float(bl.get("connect_timeout_s", 2.0)),
        ),
        ring_light=RingLightConfig(
            robot_output_pin=int(rl.get("robot_output_pin", 12)),
        ),
        pipette=PipetteConfig(
            interface=pp.get("interface", "tcp"),
            ip=pp.get("ip", ""),
            port=int(pp.get("port", 503)),
            aspirate_volume_ul=float(pp["aspirate_volume_ul"]),
            aspirate_speed_ul_per_s=float(pp["aspirate_speed_ul_per_s"]),
            dispense_volume_ul=float(pp["dispense_volume_ul"]),
            dispense_speed_ul_per_s=float(pp["dispense_speed_ul_per_s"]),
        ),
        paths=PathsConfig(
            db=pa["db"],
            plate_configs=pa["plate_configs"],
            runs_output=pa["runs_output"],
            logs=pa["logs"],
            overhead_images=pa.get("overhead_images", "output/overhead_images"),
            unprocessed_overhead_images=pa.get(
                "unprocessed_overhead_images", "output/unprocessed_overhead"
            ),
        ),
        vision=VisionConfig(
            detection_min_area_px=int(vi.get("detection_min_area_px", 50)),
            detection_max_area_px=int(vi.get("detection_max_area_px", 5000)),
            backlight_lead_time_ms=int(vi.get("backlight_lead_time_ms", 50)),
            overhead_save_images=bool(vi.get("overhead_save_images", True)),
            overhead_save_format=str(vi.get("overhead_save_format", "png")),
            eoa_save_format=str(vi.get("eoa_save_format", "tiff")),
            centroid_snap_mm=float(vi.get("centroid_snap_mm", 0.8)),
        ),
        logging=LoggingConfig(
            level=lo.get("level", "INFO"),
            retention_days=int(lo.get("retention_days", 30)),
        ),
        plate_stations=(
            PlateStationsConfig(
                right_home_u=float(ps["right_home_u"]),
                left_home_u=float(ps["left_home_u"]),
                righty_top_u=float(ps["righty_top_u"]),
                righty_bottom_u=float(ps["righty_bottom_u"]),
                lefty_top_u=float(ps["lefty_top_u"]),
                lefty_bottom_u=float(ps["lefty_bottom_u"]),
                stations={int(k): str(v["side"]) for k, v in ps["stations"].items()},
            )
            if ps is not None else None
        ),
    )
