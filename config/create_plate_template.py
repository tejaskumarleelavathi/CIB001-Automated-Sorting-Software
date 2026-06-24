"""
create_plate_template.py — Generate the plate config Excel template.

Run this once to create plate_config_template.xlsx in the same folder.
Copy that file to C:\LabAutomation\config\plate_configs\ as a starting point.

Usage:
    python config/create_plate_template.py
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).parent / "plate_config_template.xlsx"


def make_template():
    wb = openpyxl.Workbook()

    # ── Sheet 1: Data ─────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "plate_config"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    example_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        ("well",                    "Well label (e.g. A1, B12)"),
        ("x",                       "Robot X coordinate (mm)"),
        ("y",                       "Robot Y coordinate (mm)"),
        ("aspirate_volume_ul",      "Optional: aspirate volume µL (leave blank to use machine default)"),
        ("aspirate_speed_ul_per_s", "Optional: aspirate speed µL/s"),
        ("dispense_volume_ul",      "Optional: dispense volume µL"),
        ("dispense_speed_ul_per_s", "Optional: dispense speed µL/s"),
    ]

    # Header row
    for col, (name, _) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Example data rows (3 wells so operator sees the pattern)
    examples = [
        ("A1",  123.4, 56.7,  "", "", "", ""),
        ("A2",  132.4, 56.7,  "", "", "", ""),
        ("A3",  141.4, 56.7,  "", "", "", ""),
    ]
    for row_idx, values in enumerate(examples, start=2):
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val != "" else None)
            cell.fill = example_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

    # Column widths
    col_widths = [12, 14, 14, 22, 26, 22, 26]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # ── Sheet 2: Instructions ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        ("CIB Plate Config — Instructions", True, 14),
        ("", False, 11),
        ("REQUIRED COLUMNS", True, 11),
        ("  well   — Human-readable well label (e.g. A1, H12). Must match the label printed on the plate.", False, 11),
        ("  x      — Robot X coordinate in millimetres. Measured by jogging the robot to the well centre.", False, 11),
        ("  y      — Robot Y coordinate in millimetres. Same as above.", False, 11),
        ("", False, 11),
        ("OPTIONAL COLUMNS", True, 11),
        ("  aspirate_volume_ul       — How many µL to aspirate. Leave blank to use the machine default.", False, 11),
        ("  aspirate_speed_ul_per_s  — Aspiration speed. Leave blank to use the machine default.", False, 11),
        ("  dispense_volume_ul       — How many µL to dispense.", False, 11),
        ("  dispense_speed_ul_per_s  — Dispense speed.", False, 11),
        ("", False, 11),
        ("FILL ORDER", True, 11),
        ("  The order of rows in this file defines the order the robot fills the wells.", False, 11),
        ("  Row 2 = first well filled. Row 3 = second. And so on.", False, 11),
        ("  There is no separate load order file — just arrange rows in the order you want.", False, 11),
        ("", False, 11),
        ("HOW TO MEASURE X AND Y", True, 11),
        ("  1. Open RC+ software on the laptop.", False, 11),
        ("  2. Use the teach pendant or RC+ jog panel to move the robot arm.", False, 11),
        ("  3. Jog until the pipette tip is centred over the target well.", False, 11),
        ("  4. Read the X and Y values from the RC+ position display.", False, 11),
        ("  5. Enter those values into this file for that well row.", False, 11),
        ("  Do this once per well on first installation. Z height is set in the SPEL+ program.", False, 11),
        ("", False, 11),
        ("SAVING", True, 11),
        ("  Save this file as Excel (.xlsx) — do NOT save as CSV.", False, 11),
        ("  Place the file in: C:\\LabAutomation\\config\\plate_configs\\", False, 11),
        ("  The software will find it automatically — no restart needed.", False, 11),
    ]

    for row_idx, (text, bold, size) in enumerate(instructions, start=1):
        cell = ws2.cell(row=row_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=size)

    ws2.column_dimensions["A"].width = 90

    wb.save(str(OUTPUT))
    print(f"Template written to: {OUTPUT}")


if __name__ == "__main__":
    make_template()
