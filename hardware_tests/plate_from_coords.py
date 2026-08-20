"""
plate_from_coords.py — write a plate config Excel file directly from 3
already-known reference well coordinates (A1, A12, H1), no robot connection,
no interactive nudging.

Use this when you've already got A1/A12/H1 coordinates from some other
source (e.g. read off the teach pendant by hand) and just want them turned
into the full 96-well grid + Excel file. For interactively teaching those
3 points with the robot itself, use teach_plate_wells.py instead.

Same grid math as teach_plate_wells.py:
  col_vector = (A12 - A1) / (COLS - 1)   -- per-column step (A1 -> A12 direction)
  row_vector = (H1  - A1) / (ROWS - 1)   -- per-row step    (A1 -> H1  direction)
  well(r, c) = A1 + c*col_vector + r*row_vector

USAGE
  & "C:/Users/TejasKumarLeelavathi/miniconda3/envs/cib_env/python.exe" hardware_tests/plate_from_coords.py
"""

import os

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Measured reference wells (from teach pendant, re-taught 2026-08-19 —
# plate/dish position had shifted from the 2026-08-14 measurement below) ───
# (x, y) mm — U is not stored per-well, it's one fixed PLATE_U for the whole
# station (all 3 points were read at the same U, confirming that holds).
A1  = (189.955, 287.602)
A12 = (189.955, 188.602)
H1  = (252.955, 287.602)
PLATE_U = 170.011

# Previous (2026-08-14) values, kept for reference:
#   A1  = (188.551, 285.811)
#   A12 = (188.551, 186.891)
#   H1  = (251.419, 285.811)

ROWS = 8    # A-H
COLS = 12   # 1-12

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "config", "plate_configs", "test_96well.xlsx")


def compute_grid(a1: tuple, a12: tuple, h1: tuple) -> dict:
    a1 = np.array(a1, dtype=float)
    a12 = np.array(a12, dtype=float)
    h1 = np.array(h1, dtype=float)

    col_vec = (a12 - a1) / (COLS - 1)
    row_vec = (h1 - a1) / (ROWS - 1)

    print(f"  Column vector (per well, A->12 direction): {col_vec} mm  "
          f"|v|={np.linalg.norm(col_vec):.3f}mm")
    print(f"  Row vector    (per well, A->H direction):  {row_vec} mm  "
          f"|v|={np.linalg.norm(row_vec):.3f}mm")
    print("  (nominal 96-well pitch is 9.00mm — compare against the above)")

    wells = {}
    for r in range(ROWS):
        for c in range(COLS):
            pos = a1 + c * col_vec + r * row_vec
            label = f"{chr(ord('A') + r)}{c + 1}"
            wells[label] = (float(pos[0]), float(pos[1]))
    return wells


def write_excel(wells: dict, out_path: str) -> None:
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "plate_config"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["well", "x", "y", "aspirate_volume_ul", "aspirate_speed_ul_per_s",
               "dispense_volume_ul", "dispense_speed_ul_per_s"]
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    row_idx = 2
    for r in range(ROWS):
        for c in range(COLS):
            label = f"{chr(ord('A') + r)}{c + 1}"
            x, y = wells[label]
            ws.cell(row=row_idx, column=1, value=label).border = border
            ws.cell(row=row_idx, column=2, value=round(x, 3)).border = border
            ws.cell(row=row_idx, column=3, value=round(y, 3)).border = border
            for col in (4, 5, 6, 7):
                ws.cell(row=row_idx, column=col).border = border
            row_idx += 1

    col_widths = [10, 12, 12, 20, 22, 20, 22]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"\n  Saved: {out_path}  ({len(wells)} wells)")


def main() -> None:
    print("=" * 70)
    print("  Plate config from known coordinates (no robot connection)")
    print("=" * 70)
    wells = compute_grid(A1, A12, H1)

    print("\n  Computed wells (first row):")
    for c in range(COLS):
        label = f"A{c + 1}"
        wx, wy = wells[label]
        print(f"    {label:<4} ({wx:.2f}, {wy:.2f}) mm")

    write_excel(wells, OUT_PATH)

    print(f"\n  PLATE_U = {PLATE_U}")
    print(f"  Set main.py's PLATE_U constant to {PLATE_U}.")


if __name__ == "__main__":
    main()
