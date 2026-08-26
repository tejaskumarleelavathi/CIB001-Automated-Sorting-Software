"""
test_move_well_debug.py -- plain terminal version of the plate-station
test, no GUI. Connects to test_plate_station_positions.prg (must already
be running as Main -- same as the UI tool) and prints every command sent
and every response received, so we can see exactly what the robot side
echoes back for each MOVE_WELL.

Usage:
    conda run -n cib_env python hardware_tests/test_move_well_debug.py <station>

Example:
    conda run -n cib_env python hardware_tests/test_move_well_debug.py 2
"""

import socket
import sys
from pathlib import Path

sys.path.insert(0, ".")

from openpyxl import load_workbook

IP = "172.16.9.15"
PORT = 2000
XLSX = Path(__file__).parent.parent / "config" / "plate_configs" / "plate_stations.xlsx"
STATION_SIDE = {1: "right", 2: "right", 3: "left", 4: "left"}
SIDE_U = {"right": 170.065, "left": 200.023}


def send(sock: socket.socket, cmd: str) -> str:
    print(f">>> {cmd}")
    sock.sendall((cmd + "\r\n").encode("ascii"))
    buf = b""
    while not buf.endswith(b"\n"):
        chunk = sock.recv(1)
        if not chunk:
            raise ConnectionError("Robot closed the connection unexpectedly")
        buf += chunk
    resp = buf.decode("ascii", errors="replace").strip()
    print(f"<<< {resp}")
    return resp


def main() -> None:
    station = int(sys.argv[1]) if len(sys.argv) > 1 else 2
    side = STATION_SIDE[station]
    u = SIDE_U[side]
    hand = "L" if side == "left" else "R"

    wb = load_workbook(XLSX, data_only=True)
    ws = wb[f"Station {station}"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wells = [(r[0], float(r[1]), float(r[2])) for r in rows if r[0] is not None]
    well_label, x, y = wells[0]
    print(f"Station {station}: side={side}, hand={hand}, u={u}")
    print(f"First well ({well_label}) from Excel: x={x}, y={y}\n")

    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    s.settimeout(15.0)
    s.connect((IP, PORT))

    send(s, "STATUS")
    send(s, "GO_HOME_RIGHT" if side == "right" else "GO_HOME_LEFT")
    send(s, f"MOVE_WELL {x},{y},{u},{hand}")
    send(s, "BYE")
    s.close()

    print(f"\nExpected echo: OK WELL_POSITION {x},{y}")
    print("Compare the <<< MOVE_WELL response above against that expected line.")


if __name__ == "__main__":
    main()
