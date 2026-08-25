"""
test_plate_station_ui.py -- standalone PyQt test tool for verifying
4-station plate well positions and the robot's left/right home switch,
BEFORE any of this is wired into the real orchestrator cycle logic.

Talks to robot_programs/test_plate_station_positions.prg over TCP (same
port sort_server.prg uses, port 2000/device #201 -- safe since only one
of the two programs ever runs on the controller at a time). Load that
program as the controller's active Main first (see the .prg file's own
header comment), then run this.

Layout matches the OPERATOR's view standing in front of the machine:
  LEFT of operator   = Station 1, Station 2 (physically the robot's
                        RIGHT side -- Righty hand, right_home)
  CENTER              = tray (sorting dish) -- not interactive here
  RIGHT of operator  = Station 3, Station 4 (physically the robot's
                        LEFT side -- Lefty hand, left_home)

Click one or more station buttons to select them (they highlight), then
Start Test. For EACH selected station, in station-number order: go to
that station's home point first (right_home or left_home, per its side),
move the pipette above the first well, DROP (extend/confirm/retract --
no motor aspirate/dispense), then the same for the last well, before
moving on to the next selected station. Ends by returning to right_home.

Usage:
    conda run -n cib_env python hardware_tests/test_plate_station_ui.py
"""

from __future__ import annotations

import socket
import sys
from pathlib import Path

sys.path.insert(0, ".")

from openpyxl import load_workbook
from PyQt6.QtCore import QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication, QFileDialog, QGridLayout, QGroupBox, QHBoxLayout, QLabel,
    QLineEdit, QMainWindow, QPlainTextEdit, QPushButton, QVBoxLayout, QWidget,
)

DEFAULT_XLSX = Path(__file__).parent.parent / "config" / "plate_configs" / "plate_stations.xlsx"
DEFAULT_IP = "172.16.9.15"
DEFAULT_PORT = 2000

# Confirmed real-hardware values (2026-08-20/25) -- Stations 1-2 are the
# robot's right side (Righty hand, right_home); Stations 3-4 are the
# robot's left side (Lefty hand, left_home).
STATION_SIDE = {1: "right", 2: "right", 3: "left", 4: "left"}


class RobotTestClient:
    """Blocking TCP client for test_plate_station_positions.prg's protocol."""

    def __init__(self, ip: str, port: int, timeout_s: float = 15.0) -> None:
        self._sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        self._sock.settimeout(timeout_s)
        self._sock.connect((ip, port))

    def send(self, cmd: str) -> str:
        self._sock.sendall((cmd + "\r\n").encode("ascii"))
        buf = b""
        while not buf.endswith(b"\n"):
            chunk = self._sock.recv(1)
            if not chunk:
                raise ConnectionError("Robot closed the connection unexpectedly")
            buf += chunk
        resp = buf.decode("ascii", errors="replace").strip()
        if resp.startswith("ERR"):
            raise RuntimeError(f"{cmd!r} -> {resp}")
        return resp

    def close(self) -> None:
        try:
            self.send("BYE")
        except Exception:
            pass
        self._sock.close()


def load_stations(xlsx_path: Path) -> dict[int, list[tuple[str, float, float]]]:
    """Returns {station_number: [(well_label, x, y), ...]} for every 'Station N' sheet present."""
    wb = load_workbook(xlsx_path, data_only=True)
    stations: dict[int, list[tuple[str, float, float]]] = {}
    for name in wb.sheetnames:
        if not name.startswith("Station "):
            continue
        try:
            n = int(name.split(" ")[1])
        except (IndexError, ValueError):
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wells = [(r[0], float(r[1]), float(r[2])) for r in rows if r[0] is not None]
        if wells:
            stations[n] = wells
    return stations


class TestWorker(QThread):
    log = pyqtSignal(str)
    finished_ok = pyqtSignal()
    failed = pyqtSignal(str)

    def __init__(self, ip: str, port: int, selected_stations: list[int],
                 stations: dict[int, list[tuple[str, float, float]]],
                 station_u: dict[str, float]) -> None:
        super().__init__()
        self._ip = ip
        self._port = port
        self._selected = sorted(selected_stations)
        self._stations = stations
        self._station_u = station_u

    def run(self) -> None:
        client = None
        try:
            self.log.emit(f"Connecting to {self._ip}:{self._port} ...")
            client = RobotTestClient(self._ip, self._port)
            self.log.emit("Connected. " + client.send("STATUS"))

            for station in self._selected:
                wells = self._stations.get(station)
                if not wells:
                    self.log.emit(f"Station {station}: no well data loaded, skipping.")
                    continue
                side = STATION_SIDE[station]
                u = self._station_u[side]
                hand = "L" if side == "left" else "R"

                self.log.emit(f"--- Station {station} ({side} side) ---")
                self.log.emit(f"Moving to {side}_home ...")
                resp = client.send("GO_HOME_RIGHT" if side == "right" else "GO_HOME_LEFT")
                self.log.emit(resp)

                for label, (well_label, x, y) in (
                    ("first", wells[0]),
                    ("last", wells[-1]),
                ):
                    self.log.emit(f"Moving to {well_label} ({label} well) at ({x}, {y})...")
                    resp = client.send(f"MOVE_WELL {x},{y},{u},{hand}")
                    self.log.emit(resp)
                    self.log.emit(f"Dropping at {well_label}...")
                    resp = client.send("DROP")
                    self.log.emit(resp)

            self.log.emit("Returning to right_home to finish...")
            self.log.emit(client.send("GO_HOME_RIGHT"))
            self.log.emit("Done.")
            self.finished_ok.emit()
        except Exception as exc:
            self.failed.emit(str(exc))
        finally:
            if client is not None:
                client.close()


class StationButton(QPushButton):
    def __init__(self, label: str) -> None:
        super().__init__(label)
        self.setCheckable(True)
        self.setMinimumSize(140, 100)
        self._refresh_style()
        self.toggled.connect(self._refresh_style)

    def _refresh_style(self) -> None:
        if self.isChecked():
            self.setStyleSheet("background-color: #4CAF50; color: white; font-weight: bold;")
        else:
            self.setStyleSheet("background-color: #e0e0e0; color: black;")


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Plate Station Position Test")
        self._stations: dict[int, list[tuple[str, float, float]]] = {}
        self._worker: TestWorker | None = None

        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)

        # ── File / connection row ──────────────────────────────────────
        top_row = QHBoxLayout()
        self.xlsx_field = QLineEdit(str(DEFAULT_XLSX))
        browse_btn = QPushButton("Browse...")
        browse_btn.clicked.connect(self._browse_xlsx)
        load_btn = QPushButton("Load Excel")
        load_btn.clicked.connect(self._load_excel)
        top_row.addWidget(QLabel("Plate config:"))
        top_row.addWidget(self.xlsx_field, stretch=1)
        top_row.addWidget(browse_btn)
        top_row.addWidget(load_btn)
        root.addLayout(top_row)

        conn_row = QHBoxLayout()
        self.ip_field = QLineEdit(DEFAULT_IP)
        self.port_field = QLineEdit(str(DEFAULT_PORT))
        self.right_u_field = QLineEdit("170.065")
        self.left_u_field = QLineEdit("200.023")
        conn_row.addWidget(QLabel("Robot IP:"))
        conn_row.addWidget(self.ip_field)
        conn_row.addWidget(QLabel("Port:"))
        conn_row.addWidget(self.port_field)
        conn_row.addWidget(QLabel("Right-side U:"))
        conn_row.addWidget(self.right_u_field)
        conn_row.addWidget(QLabel("Left-side U:"))
        conn_row.addWidget(self.left_u_field)
        root.addLayout(conn_row)

        # ── Operator-view layout: left plates | tray | right plates ────
        layout_box = QGroupBox("Click one or more stations to select, then Start Test")
        layout_grid = QGridLayout(layout_box)

        self.station_buttons: dict[int, StationButton] = {}
        for i, station in enumerate((1, 2)):
            btn = StationButton(f"Station {station}\n(robot right / right_home)")
            self.station_buttons[station] = btn
            layout_grid.addWidget(btn, i, 0)

        tray_label = QLabel("TRAY\n(sorting dish)")
        tray_label.setStyleSheet("background-color: #ccc; padding: 20px;")
        layout_grid.addWidget(tray_label, 0, 1, 2, 1)

        for i, station in enumerate((3, 4)):
            btn = StationButton(f"Station {station}\n(robot left / left_home)")
            self.station_buttons[station] = btn
            layout_grid.addWidget(btn, i, 2)

        root.addWidget(layout_box)

        # ── Start / log ─────────────────────────────────────────────────
        self.start_btn = QPushButton("Start Test")
        self.start_btn.clicked.connect(self._start_test)
        root.addWidget(self.start_btn)

        self.log_view = QPlainTextEdit()
        self.log_view.setReadOnly(True)
        root.addWidget(self.log_view, stretch=1)

        self._load_excel()

    def _browse_xlsx(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Select plate config Excel", "", "Excel (*.xlsx)")
        if path:
            self.xlsx_field.setText(path)
            self._load_excel()

    def _load_excel(self) -> None:
        path = Path(self.xlsx_field.text())
        try:
            self._stations = load_stations(path)
            self._log(f"Loaded {path.name}: stations {sorted(self._stations.keys())}, "
                       f"well counts {[len(v) for v in self._stations.values()]}")
        except Exception as exc:
            self._log(f"Failed to load {path}: {exc}")
            self._stations = {}

    def _log(self, msg: str) -> None:
        self.log_view.appendPlainText(msg)

    def _start_test(self) -> None:
        selected = [s for s, btn in self.station_buttons.items() if btn.isChecked()]
        if not selected:
            self._log("No stations selected.")
            return
        if not self._stations:
            self._log("No plate config loaded.")
            return
        try:
            station_u = {"right": float(self.right_u_field.text()),
                         "left": float(self.left_u_field.text())}
            port = int(self.port_field.text())
        except ValueError:
            self._log("Invalid IP/port/U value.")
            return

        self.start_btn.setEnabled(False)
        self._worker = TestWorker(self.ip_field.text(), port, selected, self._stations, station_u)
        self._worker.log.connect(self._log)
        self._worker.finished_ok.connect(lambda: self.start_btn.setEnabled(True))
        self._worker.failed.connect(self._on_test_failed)
        self._worker.start()

    def _on_test_failed(self, msg: str) -> None:
        self._log(f"FAILED: {msg}")
        self.start_btn.setEnabled(True)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    win = MainWindow()
    win.resize(700, 600)
    win.show()
    sys.exit(app.exec())
